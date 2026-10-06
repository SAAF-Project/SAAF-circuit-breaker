"""Build a five-slide nontechnical SAAF explainer from one layout source.

Produces an editable PPTX, an offline HTML deck, JSON content and speaker notes.
PDF is exported from the verified HTML deck, not approximated from PowerPoint.
"""
from __future__ import annotations

import base64
import html
import json
import re
import urllib.request
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent
SUPPORT = OUT / "assets"
BG, INK, MUTED = "F6F2E8", "18323C", "536872"
ORANGE, TEAL, BORDER = "D94B27", "146B57", "D7DED7"
slides: list[dict] = []


def new_slide(title: str, eyebrow: str, subtitle: str, notes: str) -> dict:
    slide = {"title": title.replace("\n", " "), "notes": notes, "elements": []}
    slides.append(slide)
    text(slide, eyebrow, 100, 65, 1720, 40, 24, TEAL, bold=True)
    text(slide, title, 100, 142, 1720, 130, 82, INK, bold=True)
    if subtitle:
        text(slide, subtitle, 100, 282, 1720, 90, 34, MUTED)
    rect(slide, 100, 1025, 1720, 2, BORDER, radius=0)
    text(slide, "SAAF CIRCUIT BREAKER  /  JET CARE EXAMPLE", 100, 1040, 1490, 30, 20, MUTED)
    text(slide, f"{len(slides)} / 5", 1730, 1040, 90, 30, 20, MUTED, align="right")
    return slide


def text(slide, value, x, y, w, h, size=32, color=INK, bold=False, align="left"):
    slide["elements"].append({"kind": "text", "text": value, "x": x, "y": y, "w": w, "h": h,
                              "size": size, "color": color, "bold": bold, "align": align})


def rect(slide, x, y, w, h, fill="FFFFFF", radius=20, stroke=None):
    slide["elements"].append({"kind": "rect", "x": x, "y": y, "w": w, "h": h,
                              "fill": fill, "radius": radius, "stroke": stroke})


def connector(slide, x, y, w, h):
    slide["elements"].append({"kind": "connector", "x": x, "y": y, "w": w, "h": h,
                              "fill": BORDER, "radius": 0, "stroke": None})


# === SLIDE 1: THE PROBLEM IN ONE CUSTOMER STORY ===
s = new_slide("Let the AI suggest.", "THE IDEA  /  A SAFETY LAYER OUTSIDE THE AI", "",
    "Open with a familiar customer-service situation, not software terminology. An assistant wants a happy customer and a good rating. The official rules still apply. The idea is to keep those rules outside the assistant's changeable instructions. We evaluate proposed actions in a local demo; we do not send actual refunds.")
text(s, "Keep the rules in charge.", 100, 257, 1720, 130, 82, ORANGE, bold=True)
s["elements"][1]["h"] = 108
text(s, "A helpful assistant should not talk itself out of the refund policy.", 100, 427, 1720, 90, 38, MUTED)
for x, w, label, body, fill, color in [
    (100, 510, "THE CUSTOMER", "“Refund €21.35.\nI have no photo.”", "FFFFFF", INK),
    (680, 510, "THE BUSINESS RULE", "€20 maximum.\nEvidence required.", "E5EEE7", TEAL),
    (1260, 560, "THE RISK", "High customer ratings\nmust not change the rules.", "FBE5DC", ORANGE),
]:
    rect(s, x, 580, w, 245, fill)
    text(s, label, x + 30, 610, w - 60, 40, 23, color, bold=True)
    text(s, body, x + 30, 672, w - 60, 130, 38, color, bold=True)
text(s, "The assistant proposes. An independent guard can say no.", 100, 900, 1720, 65, 40, INK, bold=True)

# === SLIDE 2: WHY THERE ARE FOUR STAGES ===
s = new_slide("Four stages. Four different jobs.", "THE WORKFLOW", "Inspect + practice before launch. Guard before action. Check afterward.",
    "Use this analogy: an inspector, a practice drill, a safety barrier and a review of the receipts. They are complementary, not four versions of the same check. Forge reads the setup. Rehearsal tries test situations. Sentinel decides whether a proposed action may continue. Verify re-checks the recorded evidence afterward. Shared policy is the common thread. The prototype has separate tools; this diagram is the intended workflow, not a claim of a fully automated end-to-end orchestration pipeline.")
for i, (label, job, caption, fill, color) in enumerate([
    ("FORGE", "Inspect the\nsetup", "Is the assistant\ndesigned safely?", "FFFFFF", INK),
    ("REHEARSAL", "Practice under\npressure", "What goes wrong\nin test situations?", "E5EEE7", TEAL),
    ("SENTINEL", "Guard the\naction", "Should this proposal\ncontinue or stop?", "FBE5DC", ORANGE),
    ("VERIFY", "Check the\nrecord", "Does the saved record\nre-check correctly?", "E5EDF0", INK),
]):
    x = 100 + i * 440
    rect(s, x, 420, 390, 365, fill)
    text(s, f"0{i+1}  /  {label}", x + 28, 448, 334, 45, 24, color, bold=True)
    text(s, job, x + 28, 526, 334, 135, 45, color, bold=True)
    text(s, caption, x + 28, 685, 334, 90, 29, MUTED)
    if i < 3:
        text(s, "→", x + 392, 550, 48, 65, 44, MUTED, align="center")
text(s, "Not four versions of the same check.", 100, 869, 1720, 70, 43, INK, bold=True)
text(s, "One shared policy; different inputs, different decisions, different outputs.", 100, 945, 1720, 50, 29, MUTED)

# === SLIDE 3: STAGES 1 AND 2, WITH EXPLICIT INPUTS / OUTPUTS ===
s = new_slide("Inspect it. Then stress-test it.", "BEFORE CUSTOMER USE  /  STAGES 1 & 2", "One reads the design. The other practices difficult cases.",
    "Forge's input is the program and official rules. It flags risky patterns such as self-editing instructions, performance pressure in context and missing decision logs, then emits a risk list and policy files. These files do not themselves guarantee safety; the runtime must enforce the rules. Rehearsal takes the rules and the canonical 36-ticket scenario. In this build, rewrites and pressure events are scripted, and the 16 personas are profiles, not independent interacting agents. The output is simulated loss and a trace of what happened, useful for comparing protection on versus off.")
for x, label, color, fill, rows in [
    (100, "01 / FORGE · THE INSPECTOR", INK, "FFFFFF", [
        ("INPUT", "Agent program + official refund rules."),
        ("WHAT IT DOES", "Finds risky choices:\nchanging its rules, chasing ratings."),
        ("OUTPUT", "Risk list + rule files."),
    ]),
    (1000, "02 / REHEARSAL · THE PRACTICE ROOM", TEAL, "E5EEE7", [
        ("INPUT", "Rules + 36 test tickets."),
        ("WHAT IT DOES", "Runs a scripted pressure scenario:\nwhat if “good ratings” win?"),
        ("OUTPUT", "Simulated losses + recorded decisions."),
    ]),
]:
    rect(s, x, 394, 820, 538, fill)
    text(s, label, x + 36, 422, 748, 48, 27, color, bold=True)
    for y, (row_label, value) in zip([504, 639, 802], rows):
        text(s, row_label, x + 36, y, 748, 35, 22, MUTED, bold=True)
        text(s, value, x + 36, y + 42, 748, 112, 33 if row_label == "WHAT IT DOES" else 35, color,
             bold=row_label == "OUTPUT")
text(s, "Rehearsal is a safe practice room—not a real refund run.", 100, 960, 1720, 48, 29, MUTED)

# === SLIDE 4: STAGE 3 AND THE LIVE LOCAL INTERACTION ===
s = new_slide("Stop the action. Restore safe rules.", "DURING A PROPOSED ACTION  /  STAGE 3: SENTINEL",
    "Input: refund request + current instructions + saved lessons",
    "A tool proposal is intercepted before any action is authorized. The hard check validates money, order value and the photo flag. Instruction and memory checks look for weakened rules or lessons promoting exceptions. Real Jev judgment helps assess that state, but it cannot override a failed money/evidence guard. The local runtime can halt the proposal, record a review item, and restore its server-owned last safe instructions and memory. This is a real persisted local transition, not restoration of an external production agent. Review is a local queue, not a notification sent to a real supervisor. A Jev failure or uncertain result halts for review. No external payment is executed.")
for x, label, body, fill, color in [
    (100, "CHECK THE ACTION", "Within €20?\nWithin order value?\nPhoto marked present?", "FFFFFF", INK),
    (695, "CHECK THE INSTRUCTIONS", "Did the assistant\nweaken the original rules?", "E5EEE7", TEAL),
    (1290, "CHECK THE LESSONS", "Are its saved notes\nencouraging exceptions?", "FBE5DC", ORANGE),
]:
    rect(s, x, 422, 530, 270, fill)
    text(s, label, x + 30, 457, 470, 45, 23, color, bold=True)
    text(s, body, x + 30, 527, 470, 155, 35, color, bold=True)
for cx in [365, 960, 1555]:
    connector(s, cx, 692, 3, 48)
connector(s, 365, 738, 1193, 3)
rect(s, 620, 716, 680, 52, INK, radius=12)
text(s, "Together → one safety decision", 640, 724, 640, 40, 27, "FFFFFF", bold=True, align="center")
connector(s, 960, 768, 3, 20)
connector(s, 365, 786, 1193, 3)
for cx in [365, 960, 1555]:
    connector(s, cx, 786, 3, 14)
for x, body, fill, color in [
    (100, "Continue only\nif safe", "E5EEE7", TEAL),
    (695, "Stop → local\nhuman review", "FFFFFF", INK),
    (1290, "Bad rewrite → restore\nsafe instructions + notes", "FBE5DC", ORANGE),
]:
    rect(s, x, 800, 530, 123, fill)
    text(s, body, x + 26, 814, 478, 98, 32, color, bold=True, align="center")
text(s, "Output: decision + saved evidence. Jev is the judgment model—not the rule-maker.", 100, 955, 1720, 50, 29, MUTED)

# === SLIDE 5: STAGE 4 AND AN HONEST DEMO BOUNDARY ===
s = new_slide("Check the record—not the AI’s word.", "AFTER THE DECISION  /  STAGE 4: VERIFY",
    "A decision receipt another person can re-check—without asking an AI.",
    "Verify reads the saved decision log and approved rules, recalculates full-event hashes and parent links, and replays available guard inputs. It returns a consistency pass/fail and a readable report; changing a recorded verdict without recalculating the record fails. It cannot stop an action that already happened and does not prove a photo or a customer's story was true. Hash consistency is not a digital signature, proof of origin or a compliance certificate. Today the web demo has real Jev calls, hard guards, persisted local rollback and offline checks. The rehearsal remains scripted; signed transitions, external agent/container control and actual refunds are not connected. The workflow handoffs are not fully automated. End by inviting the organiser to see ticket 36 and the decision receipt.")
for x, label, body, fill, color in [
    (100, "INPUT", "Saved decision log\n+ approved rules", "FFFFFF", INK),
    (695, "CHECK", "Recheck recorded details\nagainst the rules", "E5EEE7", TEAL),
    (1290, "OUTPUT", "Consistent / mismatch\n+ a readable report", "E5EDF0", INK),
]:
    rect(s, x, 412, 530, 295, fill)
    text(s, label, x + 30, 449, 470, 45, 25, color, bold=True)
    text(s, body, x + 30, 526, 470, 146, 37, color, bold=True)
for x in [637, 1232]:
    text(s, "→", x, 552, 50, 65, 44, MUTED, align="center")
text(s, "If saved details do not match → the check fails.", 100, 773, 1720, 65, 37, INK, bold=True)
text(s, "Checkable evidence—not a digital signature or compliance certificate.", 100, 849, 1720, 54, 29, MUTED)
rect(s, 100, 934, 1720, 80, INK)
text(s, "Demo: real Jev + local controls; rehearsal is scripted.\nNot fully automated. No real refunds, signed records or certified compliance.",
     130, 941, 1660, 68, 26, "FFFFFF")

# === VALIDATE THE SOURCE BEFORE EXPORT ===
assert len(slides) == 5
for i, slide in enumerate(slides, 1):
    assert slide["notes"]
    for element in slide["elements"]:
        assert 0 <= element["x"] and 0 <= element["y"]
        assert element["x"] + element["w"] <= 1920, (i, element)
        assert element["y"] + element["h"] <= 1080, (i, element)

# === EDITABLE POWERPOINT: NATIVE TEXT AND VECTOR PANELS ===
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(20), Inches(11.25)
prs.core_properties.title = "SAAF Circuit Breaker — a plain-language explainer"
prs.core_properties.subject = "Five stages-and-interaction slides for hackathon organisers"
prs.core_properties.author = "SAAF prototype team"
for source in slides:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    for e in source["elements"]:
        xywh = [Inches(e[k] / 96) for k in ["x", "y", "w", "h"]]
        if e["kind"] in {"rect", "connector"}:
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE, *xywh)
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(e["fill"])
            shape.line.fill.background()
            if e["radius"]:
                shape.adjustments[0] = 0.07
        else:
            shape = slide.shapes.add_textbox(*xywh)
            frame = shape.text_frame
            frame.clear()
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.TOP
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            for line_index, line in enumerate(e["text"].split("\n")):
                paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.name = "DM Sans"
                paragraph.font.size = Pt(e["size"] * 0.75)
                paragraph.font.bold = e["bold"]
                paragraph.font.color.rgb = RGBColor.from_string(e["color"])
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.14
                paragraph.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[e["align"]]
    slide.notes_slide.notes_text_frame.text = source["notes"]
prs.save(OUT / "SAAF-plain-language.pptx")

# === SELF-CONTAINED HTML: EMBED GOOGLE FONT DATA, NO LIVE CDN NEEDED ===
request = urllib.request.Request("https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;600;700&display=swap", headers={"User-Agent": "Mozilla/5.0"})
font_css = urllib.request.urlopen(request, timeout=20).read().decode()
for url in set(re.findall(r"url\((https://[^)]+)\)", font_css)):
    blob = urllib.request.urlopen(url, timeout=20).read()
    font_css = font_css.replace(url, "data:font/woff2;base64," + base64.b64encode(blob).decode())
base_css = (SUPPORT / "viewport-base.css").read_text()
sections = []
for i, slide in enumerate(slides):
    pieces = []
    for e in slide["elements"]:
        position = ";".join(f"{k}:{e[v]}px" for k, v in [("left", "x"), ("top", "y"), ("width", "w"), ("height", "h")])
        if e["kind"] in {"rect", "connector"}:
            css_class = "panel" if e["kind"] == "rect" else "connector"
            pieces.append(f'<div class="{css_class}" aria-hidden="true" style="{position};background:#{e["fill"]};border-radius:{e["radius"]}px"></div>')
        else:
            pieces.append(f'<div class="copy" data-editable="true" style="{position};font-size:{e["size"]}px;color:#{e["color"]};font-weight:{700 if e["bold"] else 400};text-align:{e["align"]}">{html.escape(e["text"])}</div>')
    sections.append(f'<section class="slide {"active visible" if i == 0 else ""}" aria-label="{html.escape(slide["title"])}" data-slide="{i+1}">' + "".join(pieces) + f'<aside class="speaker-notes" hidden>{html.escape(slide["notes"])}</aside></section>')
custom_css = '''
/* === THEME / AUTHORING GRID === */
:root { --stage-bg:#18323c; --slide-bg:#F6F2E8; }
* { box-sizing:border-box; }
body { font-family:'DM Sans',sans-serif; }
.panel,.connector,.copy { position:absolute; }
.copy { white-space:pre-wrap; line-height:1.14; overflow-wrap:normal; }
.slide { transition:opacity .2s ease; }
.deck-controls { display:flex; gap:12px; padding:8px 14px; border-radius:30px; background:#18323ceb; color:#fff; font-size:14px; }
button { font:inherit; cursor:pointer; border:0; border-radius:18px; padding:6px 14px; color:#18323c; background:#f6f2e8; }
.deck-controls a { color:#fff; padding:6px; }
.edit-hotzone { position:fixed; top:0; left:0; width:60px; height:60px; z-index:2000; }
#editToggle { position:fixed; top:15px; left:15px; z-index:2001; opacity:0; pointer-events:none; }
#editToggle.show,#editToggle.active { opacity:1; pointer-events:auto; }
[contenteditable=true] { outline:2px dashed #d94b27; }
.capture .deck-controls,.capture .edit-hotzone,.capture #editToggle { display:none; }
.capture * { transition:none!important; animation:none!important; }
@page { size:20in 11.25in; margin:0; }
@media print { * { -webkit-print-color-adjust:exact; print-color-adjust:exact; } #editToggle,.edit-hotzone { display:none!important; } }
'''
controller = '''
/* === FIXED-STAGE CONTROLLER / ACCESSIBLE NAVIGATION === */
class SlidePresentation {
  constructor(){this.slides=[...document.querySelectorAll('.slide')];this.currentSlide=0;this.stage=document.querySelector('.deck-stage');this.isEditing=false;this.scale();addEventListener('resize',()=>this.scale());new ResizeObserver(()=>this.scale()).observe(document.documentElement);window.visualViewport?.addEventListener('resize',()=>this.scale());this.showSlide(Number(location.hash.slice(1)||1)-1);}
  scale(){const s=Math.min(innerWidth/1920,innerHeight/1080);this.stage.style.transform=`translate(${(innerWidth-1920*s)/2}px,${(innerHeight-1080*s)/2}px) scale(${s})`;}
  showSlide(n){this.currentSlide=Math.max(0,Math.min(this.slides.length-1,n));this.slides.forEach((slide,i)=>{slide.classList.toggle('active',i===this.currentSlide);slide.classList.toggle('visible',i===this.currentSlide);slide.setAttribute('aria-hidden',String(i!==this.currentSlide));});document.querySelector('#counter').textContent=`${this.currentSlide+1} / ${this.slides.length}`;history.replaceState(null,'',`#${this.currentSlide+1}`);}
  toggleEdit(){this.isEditing=!this.isEditing;document.querySelectorAll('[data-editable]').forEach(e=>e.contentEditable=String(this.isEditing));document.querySelector('#editToggle').classList.toggle('active',this.isEditing);}
}
const deck=new SlidePresentation();window.deck=deck;window.goSlide=n=>deck.showSlide(n);
document.querySelector('#prev').onclick=()=>deck.showSlide(deck.currentSlide-1);
document.querySelector('#next').onclick=()=>deck.showSlide(deck.currentSlide+1);
document.querySelector('#editToggle').onclick=()=>deck.toggleEdit();
addEventListener('keydown',e=>{if(e.target.isContentEditable)return;if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();deck.showSlide(deck.currentSlide+1);}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();deck.showSlide(deck.currentSlide-1);}if(e.key==='Home')deck.showSlide(0);if(e.key==='End')deck.showSlide(deck.slides.length-1);if(e.key.toLowerCase()==='e')deck.toggleEdit();});
let touchX=0;addEventListener('touchstart',e=>touchX=e.changedTouches[0].clientX,{passive:true});addEventListener('touchend',e=>{if(deck.isEditing)return;const d=e.changedTouches[0].clientX-touchX;if(Math.abs(d)>45)deck.showSlide(deck.currentSlide+(d<0?1:-1));},{passive:true});
let lastWheel=0;addEventListener('wheel',e=>{if(!deck.isEditing&&Date.now()-lastWheel>400&&Math.abs(e.deltaY)>12){deck.showSlide(deck.currentSlide+(e.deltaY>0?1:-1));lastWheel=Date.now();}},{passive:true});
/* === INLINE EDITING: LOCAL SAVE, NEVER CHANGES THE APPLICATION === */
const editable=[...document.querySelectorAll('[data-editable]')];
try{const saved=JSON.parse(localStorage.getItem('saaf-organiser-copy-v1')||'null');if(saved?.length===editable.length)editable.forEach((e,i)=>e.innerText=saved[i]);}catch{}
editable.forEach(e=>e.addEventListener('input',()=>{try{localStorage.setItem('saaf-organiser-copy-v1',JSON.stringify(editable.map(x=>x.innerText)));}catch{}}));
const zone=document.querySelector('.edit-hotzone'),toggle=document.querySelector('#editToggle');let hide;
for(const e of [zone,toggle]){e.addEventListener('mouseenter',()=>{clearTimeout(hide);toggle.classList.add('show');});e.addEventListener('mouseleave',()=>hide=setTimeout(()=>{if(!deck.isEditing)toggle.classList.remove('show');},400));}
zone.onclick=()=>deck.toggleEdit();
addEventListener('keydown',e=>{if((e.ctrlKey||e.metaKey)&&e.key.toLowerCase()==='s'&&deck.isEditing){e.preventDefault();document.querySelectorAll('[data-editable]').forEach(x=>x.contentEditable='false');const data='<!doctype html>'+document.documentElement.outerHTML;const a=document.createElement('a');a.href=URL.createObjectURL(new Blob([data],{type:'text/html'}));a.download='SAAF-plain-language-edited.html';a.click();URL.revokeObjectURL(a.href);document.querySelectorAll('[data-editable]').forEach(x=>x.contentEditable='true');}});
'''
page = f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>SAAF — five slides in plain language</title><style>{font_css}\n{base_css}\n{custom_css}</style></head><body><div class="deck-viewport"><main class="deck-stage">{"".join(sections)}</main></div><nav class="deck-controls" aria-label="Presentation controls"><button id="prev" aria-label="Previous slide">←</button><span id="counter">1 / 5</span><button id="next" aria-label="Next slide">→</button><a href="SAAF-plain-language.pdf">PDF</a><a href="SAAF-plain-language.pptx">PowerPoint</a></nav><div class="edit-hotzone" title="Edit text"></div><button id="editToggle" aria-label="Toggle text editing">Edit text · E</button><script>{controller}</script></body></html>'''
(OUT / "SAAF-plain-language.html").write_text(page, encoding="utf-8")
(OUT / "deck-source.json").write_text(json.dumps(slides, indent=2, ensure_ascii=False), encoding="utf-8")
(OUT / "speaker-notes.md").write_text("# SAAF organiser explainer — speaking notes\n\n" + "\n\n".join(f"## {i}. {s['title']}\n\n{s['notes']}" for i, s in enumerate(slides, 1)), encoding="utf-8")
readback = Presentation(OUT / "SAAF-plain-language.pptx")
assert len(readback.slides) == 5
for native, source in zip(readback.slides, slides):
    actual = [shape.text for shape in native.shapes if shape.has_text_frame and shape.text]
    expected = [e["text"] for e in source["elements"] if e["kind"] == "text"]
    assert actual == expected
    assert native.notes_slide.notes_text_frame.text == source["notes"]
print(json.dumps({"slides": len(readback.slides), "editable_text_verified": True, "speaker_notes_verified": True,
                  "files": [str(OUT / name) for name in ["SAAF-plain-language.pptx", "SAAF-plain-language.html", "speaker-notes.md"]]}, indent=2))
