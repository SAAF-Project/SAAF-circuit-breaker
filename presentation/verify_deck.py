"""Verify PPTX/PDF content and render every PDF page for review."""
import json
from pathlib import Path

import pypdfium2 as pdfium
from PIL import Image, ImageDraw
from pypdf import PdfReader
from pptx import Presentation

root = Path(__file__).resolve().parent
pdf = root / "SAAF-plain-language.pdf"
reader = PdfReader(pdf)
prs = Presentation(root / "SAAF-plain-language.pptx")
assert len(reader.pages) == len(prs.slides) == 5
expected = ["Let the AI suggest.", "Four stages. Four different jobs.", "Inspect it. Then stress-test it.",
            "Stop the action. Restore safe rules.", "Check the record—not the AI’s word."]
records = []
for i, (page, title) in enumerate(zip(reader.pages, expected), 1):
    text = page.extract_text()
    assert title in text, (i, title, text)
    assert f"{i} / 5" in text
    width, height = float(page.mediabox.width), float(page.mediabox.height)
    assert abs(width / height - 16 / 9) < 0.001
    assert prs.slides[i-1].notes_slide.notes_text_frame.text.strip()
    records.append({"page": i, "title": title, "characters": len(text), "width": width, "height": height})
preview = root / "preview"
preview.mkdir(exist_ok=True)
doc = pdfium.PdfDocument(str(pdf))
images = []
for i in range(len(doc)):
    page = doc[i]
    bitmap = page.render(scale=1280 / page.get_width())
    image = bitmap.to_pil().convert("RGB")
    image.save(preview / f"pdf-slide-{i+1}.png")
    images.append(image)
    bitmap.close()
    page.close()
doc.close()
contact = Image.new("RGB", (2560, 2160), "#18323C")
for i, image in enumerate(images):
    contact.paste(image, ((i % 2) * 1280, (i // 2) * 720))
contact.save(preview / "all-slides.png")
(root / "verification.json").write_text(json.dumps({"pptx_slides": len(prs.slides), "pdf_pages": len(reader.pages),
    "page_checks": records, "speaker_notes": True, "pdf_rendered_pages": len(images)}, indent=2), encoding="utf-8")
print(json.dumps({"pptx_slides": len(prs.slides), "pdf_pages": len(reader.pages), "pdf_rendered_pages": len(images),
                  "contact_sheet": str(preview / "all-slides.png")}, indent=2))
